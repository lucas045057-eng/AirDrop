"""项目账号资料：加密保存、遮蔽预览、逐项冲突选择与幂等确认。

资料按 (项目, 脚本包, 钱包) 维度保存，同一钱包可以在不同项目或脚本里使用不同
代理、邀请码等参数；敏感值只以加密形式落库，预览与列表只返回遮蔽结果。
"""
import base64
import csv
import hashlib
import io
import json
import re
import secrets
import time
import zipfile
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import func, select
from backend.config import settings
from backend.models import AccountImportDraft, Artifact, AuditLog, Project, ProjectAccount, ScriptPackage, Wallet, WalletSecret, uid
from backend.storage import artifact_lock, read_blob, write_blob
from backend.vault import decrypt_wallet, encrypt_wallet, identity, key
from adapters.wallets import validate_address
from backend.uploads import redact
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

MAX_ROWS = 1000
MAX_COLUMNS = 40
DRAFT_TTL = 3600
IDENTITY_FIELDS = ('address', 'private_key')
SENSITIVE_FIELD_PATTERN = re.compile(
    r'(?<![a-z0-9])(?:private[\W_]*key|proxy|api[\W_]*key|access[\W_]*key|token|secret|'
    r'password|passwd|cookie|mnemonic|seed|authorization|credential|auth|bearer|session[\W_]*key)'
    r'(?![a-z0-9])', re.I)
FORMATS = {'.csv': 'CSV', '.json': 'JSON', '.xlsx': 'XLSX', '.txt': 'TXT'}


def mask(value: str) -> str:
    """只显示是否存在与首尾少量字符，不回传完整值。"""
    text = str(value)
    if not text:
        return ''
    if len(text) <= 6:
        return '*' * len(text)
    return text[:3] + '*' * 4 + text[-3:]


def is_sensitive(name: str) -> bool:
    normalized = re.sub(r'(?<=[a-z0-9])(?=[A-Z])', '_', name).lower()
    return bool(SENSITIVE_FIELD_PATTERN.search(normalized))


def bounded_rows(items) -> list[dict]:
    rows = []
    for row in items:
        if len(rows) >= MAX_ROWS:
            raise ValueError('资料数量应为 1–1000 条')
        rows.append(row)
    return rows


def json_rows(text: str) -> list[dict]:
    """Decode an object array incrementally so oversized files cannot expand unboundedly."""
    decoder = json.JSONDecoder()

    def whitespace(index):
        while index < len(text) and text[index].isspace():
            index += 1
        return index

    index = whitespace(0)
    if index >= len(text) or text[index] != '[':
        raise ValueError('JSON 资料必须是对象数组')
    index = whitespace(index + 1)
    rows = []
    if index < len(text) and text[index] == ']':
        index += 1
    else:
        while True:
            if len(rows) >= MAX_ROWS:
                raise ValueError('资料数量应为 1–1000 条')
            value, index = decoder.raw_decode(text, index)
            if not isinstance(value, dict):
                raise ValueError('JSON 资料必须是对象数组')
            rows.append(value)
            index = whitespace(index)
            if index < len(text) and text[index] == ',':
                index = whitespace(index + 1)
                continue
            if index < len(text) and text[index] == ']':
                index += 1
                break
            raise json.JSONDecodeError('期望数组分隔符', text, index)
    index = whitespace(index)
    if index != len(text):
        raise json.JSONDecodeError('数组后存在多余内容', text, index)
    return rows


def account_key(project_id: str, package_id: str, wallet_id: str) -> bytes:
    return f'airdrop-v2-account:{project_id}:{package_id}:{wallet_id}'.encode()


def seal(project_id, package_id, wallet_id, values: dict) -> str:
    nonce = secrets.token_bytes(12)
    ciphertext = AESGCM(key()).encrypt(nonce, json.dumps(values, sort_keys=True).encode(), account_key(project_id, package_id, wallet_id))
    return base64.b64encode(nonce + ciphertext).decode()


def open_config(account: ProjectAccount) -> dict:
    if not account.encrypted_config:
        return {}
    try:
        raw = base64.b64decode(account.encrypted_config, validate=True)
        return json.loads(AESGCM(key()).decrypt(raw[:12], raw[12:], account_key(account.project_id, account.package_id, account.wallet_id)))
    except Exception:
        raise HTTPException(409, '账号加密资料校验失败，请重新导入该钱包的资料') from None


def read_table(filename: str, data: bytes) -> tuple[str, list[dict]]:
    """按文件后缀解析表格；不执行任何内容，只读取单元格。"""
    suffix = '.' + filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    format = FORMATS.get(suffix)
    if format is None:
        raise ValueError('仅支持 CSV、JSON、XLSX、TXT 资料文件')
    try:
        if format == 'CSV':
            rows = bounded_rows(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
        elif format == 'JSON':
            rows = json_rows(data.decode('utf-8-sig'))
        elif format == 'XLSX':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 100 or sum(item.file_size for item in archive.infolist()) > 10 * 1024 * 1024:
                    raise ValueError('Excel 解压体积过大')
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            sheet = workbook.active
            if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                raise ValueError('资料行数超过 1000')
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else '' for value in next(iterator)]
            if len(set(headers)) != len(headers) or '' in headers:
                raise ValueError('Excel 列名重复或为空')
            rows = bounded_rows(dict(zip(headers, values)) for values in iterator
                                if any(value is not None for value in values))
            workbook.close()
        else:
            lines = io.StringIO(data.decode('utf-8-sig'))
            rows = bounded_rows({'value': line.strip()} for line in lines if line.strip())
    except (UnicodeError, json.JSONDecodeError, zipfile.BadZipFile, OSError, StopIteration) as exc:
        raise ValueError('资料文件损坏或格式不符') from None
    if not rows:
        raise ValueError('资料数量应为 1–1000 条')
    if any(len(row) > MAX_COLUMNS for row in rows):
        raise ValueError('资料列数过多')
    return format, rows


def columns_of(rows: list[dict]) -> list[str]:
    names = []
    for row in rows:
        for name in row:
            if name not in names:
                names.append(str(name))
    return names


def validate_mapping(mapping: dict, columns: list[str]) -> dict:
    if not isinstance(mapping, dict) or not mapping or len(mapping) > MAX_COLUMNS:
        raise ValueError('字段映射必须非空且不超过 40 项')
    clean = {}
    for target, source in mapping.items():
        if not isinstance(target, str) or not isinstance(source, str) or not target or not source:
            raise ValueError('字段映射格式无效')
        if source not in columns:
            raise ValueError('资料文件中没有列：' + source)
        clean[target] = source
    if 'address' not in clean and 'private_key' not in clean:
        raise ValueError('字段映射必须包含 address 或 private_key 中的至少一项')
    return clean


def load_rows(db, project_id, package_id, data: dict, user_id):
    """解析资料并存在有期限的加密草稿中；原始资料不落明文。"""
    filename = str(data.get('filename') or 'accounts.csv')[:200]
    body = data.get('content')
    if not isinstance(body, (bytes, bytearray)) or not body:
        raise ValueError('请上传资料文件')
    if len(body) > 4 * 1024 * 1024:
        raise ValueError('资料文件超过 4 MiB')
    chain = str(data.get('chain') or '')
    if chain not in {'EVM', 'SOLANA'}:
        raise ValueError('必须选择钱包链类型')
    format, rows = read_table(filename, bytes(body))
    mapping = validate_mapping(data.get('field_mapping'), columns_of(rows))
    path, digest = write_blob(bytes(body))
    artifact = Artifact(path=path, sha256=digest, kind='account-import')
    db.add(artifact)
    db.flush()
    draft = AccountImportDraft(project_id=project_id, package_id=package_id, artifact_id=artifact.id,
                               file_sha256=digest, filename=filename, field_mapping=mapping,
                               preview={'chain': chain, 'format': format}, preview_sha256=digest,
                               created_by=user_id, expires_at=time.time() + DRAFT_TTL)
    db.add(draft)
    db.flush()
    preview = build_preview(db, draft, chain, rows)
    draft.preview = preview
    draft.preview_sha256 = hashlib.sha256(json.dumps(preview['entries'], sort_keys=True).encode()).hexdigest()
    db.add(AuditLog(user_id=user_id, action='account.import.draft', target=draft.id,
                    detail={'filename': filename, 'rows': len(rows), 'mapping': sorted(mapping)}))
    return draft


def row_values(draft: AccountImportDraft, row: dict) -> dict:
    return {target: ('' if row.get(source) is None else str(row[source]).strip()) for target, source in draft.field_mapping.items()}


def build_preview(db, draft: AccountImportDraft, chain: str, rows: list[dict]):
    """按新增、复用、冲突、错误分组；敏感值只给出遮蔽结果。"""
    entries, errors, seen = [], [], set()
    for position, row in enumerate(rows, 1):
        values = row_values(draft, row)
        try:
            private_key = values.get('private_key', '')
            address = values.get('address', '')
            if private_key:
                derived = identity(chain, private_key)[0]
                if address and derived != address:
                    raise ValueError('地址与私钥不一致')
                address = derived
            if not address:
                raise ValueError('缺少地址或私钥')
            validate_address(chain, address)
        except ValueError as exc:
            errors.append({'row': position, 'reason': str(exc)[:200]})
            continue
        if address in seen:
            errors.append({'row': position, 'reason': '资料中地址重复'})
            continue
        seen.add(address)
        wallet = db.scalar(select(Wallet).where(Wallet.chain == chain, Wallet.address == address))
        account = db.scalar(select(ProjectAccount).where(ProjectAccount.project_id == draft.project_id,
                                                         ProjectAccount.package_id == draft.package_id,
                                                         ProjectAccount.wallet_id == wallet.id)) if wallet else None
        existing = {**account.public_config, **open_config(account)} if account else {}
        conflicts = []
        for target, value in values.items():
            if target == 'private_key':
                continue
            before = str(existing.get(target, ''))
            if before and value and before != value:
                conflicts.append({'field': target, 'current': mask(before), 'incoming': mask(value)})
        entries.append({'address': address, 'wallet': wallet.id if wallet else None,
                        'account_id': account.id if account else None, 'revision': account.revision if account else None,
                        'group': 'conflict' if conflicts else ('reuse' if account else 'new'),
                        'conflicts': conflicts,
                        'values': {target: mask(value) if is_sensitive(target) else value for target, value in sorted(values.items()) if target != 'private_key'},
                        'has_private_key': bool(private_key)})
    counts = {group: sum(1 for entry in entries if entry['group'] == group) for group in ('new', 'reuse', 'conflict')}
    counts['error'] = len(errors)
    return {'chain': chain, 'format': draft.preview.get('format'), 'counts': counts, 'entries': entries, 'errors': errors,
            'mapping': draft.field_mapping, 'file_sha256': draft.file_sha256,
            'note': '空值表示未提供，不会清空已有值；清空必须显式选择'}


def decide(entry: dict, decisions: dict) -> str:
    """冲突处理必须逐项给出；非冲突项默认按覆盖处理，空值不会清空已有值。"""
    choice = decisions.get(entry['address'])
    allowed = {'overwrite', 'skip', 'clear'}
    if choice is None:
        if entry['group'] == 'conflict':
            raise HTTPException(409, '冲突项必须逐项选择处理方式：' + mask(entry['address']))
        return 'overwrite'
    if choice not in allowed:
        raise HTTPException(422, '冲突处理选择无效：' + str(choice)[:20])
    return choice


def stored_rows(db, draft: AccountImportDraft):
    """从加密草稿重新读取原始资料；预览只保存遮蔽结果，确认时才解密使用。"""
    artifact = db.get(Artifact, draft.artifact_id)
    data = read_blob(artifact.path, artifact.sha256)
    if hashlib.sha256(data).hexdigest() != draft.file_sha256:
        raise HTTPException(409, '导入资料工件已改变，请重新上传')
    format, rows = read_table(draft.filename, data)
    return format, [row_values(draft, row) for row in rows]


def commit(db, draft: AccountImportDraft, preview_sha256: str, decisions: dict, user_id):
    """一次事务完成导入；重复确认返回同一结果。"""
    if draft.consumed_at:
        return draft.result
    if draft.expires_at < time.time():
        raise HTTPException(409, '导入草稿已过期，请重新上传并预览')
    if preview_sha256 != draft.preview_sha256:
        raise HTTPException(409, '预览摘要已改变，请重新预览后再确认')
    chain = draft.preview['chain']
    _, parsed = stored_rows(db, draft)
    by_address = {}
    for values in parsed:
        private_key, address = values.get('private_key', ''), values.get('address', '')
        if private_key:
            address = identity(chain, private_key)[0]
        if address:
            by_address[address] = values
    created = updated = skipped = 0
    wallets_created = 0
    for entry in draft.preview['entries']:
        choice = decide(entry, decisions)
        if choice == 'skip':
            skipped += 1
            continue
        wallet = db.get(Wallet, entry['wallet']) if entry['wallet'] else None
        if wallet is None:
            wallet = db.scalar(select(Wallet).where(Wallet.chain == chain, Wallet.address == entry['address']))
        account = db.scalar(select(ProjectAccount).where(ProjectAccount.project_id == draft.project_id,
                                                        ProjectAccount.package_id == draft.package_id,
                                                        ProjectAccount.wallet_id == wallet.id)) if wallet else None
        current_revision = account.revision if account else None
        if current_revision != entry['revision']:
            raise HTTPException(409, '账号资料在预览之后被修改（' + mask(entry['address']) + '），请重新预览')
        values = by_address.get(entry['address'])
        if values is None:
            raise HTTPException(409, '导入资料与预览不一致（' + mask(entry['address']) + '），请重新预览')
        if wallet is not None and entry['revision'] is None and db.get(WalletSecret, wallet.id) is None and not values.get('private_key'):
            raise HTTPException(409, '钱包保险库缺少私钥（' + mask(entry['address']) + '），请在资料中提供私钥后重新导入')
        if wallet is None:
            # 新地址必须有私钥才能建立钱包；仅凭文本地址不新建无法核对的钱包。
            if not values.get('private_key'):
                raise HTTPException(409, '新钱包必须提供私钥以核验地址：' + mask(entry['address']))
            wallet = Wallet(chain=chain, address=entry['address'], name=str(values.get('name') or '')[:100], is_test=False)
            db.add(wallet)
            db.flush()
            wallets_created += 1
        if values.get('private_key'):
            secret = db.get(WalletSecret, wallet.id)
            if secret is None:
                encrypted = encrypt_wallet(wallet, values['private_key'])
                db.add(WalletSecret(wallet_id=wallet.id, ciphertext=encrypted))
            elif decrypt_wallet(wallet, secret.ciphertext) != values['private_key']:
                raise HTTPException(409, '钱包保险库中的私钥与导入资料不一致：' + mask(entry['address']))
        merged = {**(account.public_config or {}), **open_config(account)} if account else {}
        for target, value in values.items():
            if choice == 'clear':
                merged.pop(target, None)
            elif value == '':
                continue
            else:
                merged[target] = value
        public = {name: value for name, value in merged.items() if not is_sensitive(name)}
        encrypted = {name: value for name, value in merged.items() if is_sensitive(name)}
        if account is None:
            account = ProjectAccount(project_id=draft.project_id, package_id=draft.package_id, wallet_id=wallet.id,
                                     public_config={}, encrypted_config='', revision=1, source=draft.filename)
            db.add(account)
            created += 1
        else:
            updated += 1
            account.revision = (account.revision or 0) + 1
        account.public_config = public
        account.encrypted_config = seal(draft.project_id, draft.package_id, wallet.id, encrypted) if encrypted else ''
        account.source = draft.filename
        account.updated_at = time.time()
    draft.consumed_at = time.time()
    draft.result = {'created': created, 'updated': updated, 'skipped': skipped, 'wallets_created': wallets_created,
                    'preview_sha256': draft.preview_sha256}
    db.add(AuditLog(user_id=user_id, action='account.import.commit', target=draft.id, detail=draft.result))
    return draft.result


def account_view(account: ProjectAccount, wallet: Wallet):
    hidden = {name: mask(value) for name, value in sorted(open_config(account).items())}
    return {'id': account.id, 'project_id': account.project_id, 'package_id': account.package_id,
            'wallet_id': account.wallet_id, 'address': wallet.address, 'name': wallet.name, 'chain': wallet.chain,
            'public_config': account.public_config, 'secured_config': hidden,
            'has_secrets': bool(hidden), 'revision': account.revision, 'source': account.source,
            'updated_at': account.updated_at}


def rows_for(db, project_id, package_id, wallet_ids):
    """运行输入需要的账号资料；只读取本次选中的钱包。"""
    rows = {}
    for account in db.scalars(select(ProjectAccount).where(ProjectAccount.project_id == project_id,
                                                          ProjectAccount.package_id == package_id,
                                                          ProjectAccount.wallet_id.in_(list(wallet_ids)))):
        rows[account.wallet_id] = {**account.public_config, **open_config(account)}
    return rows
