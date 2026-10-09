import csv
import io
import json
import zipfile
from pathlib import Path
from openpyxl import Workbook,load_workbook


def validate_address(chain,address):
    if chain=='EVM':
        from eth_utils import is_checksum_address
        if not is_checksum_address(address):raise ValueError('EVM 地址格式或校验和错误')
    elif chain=='SOLANA':
        from solders.pubkey import Pubkey
        try:Pubkey.from_string(address)
        except Exception:raise ValueError('Solana 地址格式错误') from None
    else:raise ValueError('链类型不支持')


def encode(records,format,mapping):
    if not mapping:raise ValueError('必须提供明确字段映射')
    rows=[]
    for record in records:
        if any(field not in record for field in mapping.values()):raise ValueError('钱包记录缺少必要字段')
        rows.append({target:record[source] for target,source in mapping.items()})
    if format=='JSON':return json.dumps(rows,ensure_ascii=False).encode()
    if format=='TXT':
        if len(mapping)!=1:raise ValueError('TXT 一行一个字段，必须只映射一个字段')
        return ('\n'.join(str(next(iter(row.values()))) for row in rows)+'\n').encode()
    if format=='CSV':
        output=io.StringIO(newline='');writer=csv.DictWriter(output,fieldnames=list(mapping))
        writer.writeheader();writer.writerows(rows)
        return output.getvalue().encode('utf-8-sig')
    if format=='XLSX':
        workbook=Workbook();sheet=workbook.active;sheet.append(list(mapping))
        for row in rows:sheet.append(list(row.values()))
        output=io.BytesIO();workbook.save(output);return output.getvalue()
    raise ValueError('不支持的钱包格式')


def decode(data,format,mapping,chain):
    if len(data)>2*1024*1024:raise ValueError('钱包文件超过 2 MiB')
    try:
        if format=='JSON':
            rows=json.loads(data)
            if not isinstance(rows,list):raise ValueError()
        elif format=='TXT':
            if len(mapping)!=1:raise ValueError('TXT 必须只映射一个字段')
            name=next(iter(mapping));rows=[{name:v.strip()} for v in data.decode('utf-8-sig').splitlines() if v.strip()]
        elif format=='CSV':
            rows=list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
        elif format=='XLSX':
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist())>100 or sum(i.file_size for i in archive.infolist())>10*1024*1024:raise ValueError('Excel 解压体积过大')
            workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
            sheet=workbook.active
            if sheet.max_row and sheet.max_row>1001:raise ValueError('Excel 行数过多')
            iterator=sheet.iter_rows(values_only=True);headers=next(iterator)
            if len(set(headers))!=len(headers):raise ValueError('Excel 列名重复')
            rows=[dict(zip(headers,r)) for r in iterator if any(v is not None for v in r)]
            workbook.close()
        else:raise ValueError('不支持的钱包格式')
        if not rows or len(rows)>1000:raise ValueError('钱包数量应为 1–1000')
        records=[]
        for row in rows:
            if not isinstance(row,dict) or any(name not in row or row[name] is None for name in mapping):
                raise ValueError('字段缺失，请显式指定字段映射')
            result={target:str(row[source]) for source,target in mapping.items()}
            if 'chain' in result and result['chain']!=chain:raise ValueError('不同链钱包不可混用')
            if 'address' in result:validate_address(chain,result['address'])
            records.append(result)
        return records
    except (UnicodeError,json.JSONDecodeError,zipfile.BadZipFile,StopIteration,KeyError,TypeError):
        raise ValueError('钱包文件损坏或格式不符') from None


def load_file(path,format,mapping,chain):
    if not Path(path).is_file():raise ValueError('钱包路径不存在')
    return decode(Path(path).read_bytes(),format,mapping,chain)
